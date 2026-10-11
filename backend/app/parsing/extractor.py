import io
import re
import os
import logging
from typing import List, Dict, Any, Tuple, Optional
import pymupdf  # PyMuPDF
import pdfplumber
from PIL import Image
import pytesseract
import docx
from pptx import Presentation

logger = logging.getLogger(__name__)

class PageExtractionResult:
    def __init__(self, page_number: int, text: str, is_scanned: bool = False, warning: Optional[str] = None):
        self.page_number = page_number
        self.text = text
        self.is_scanned = is_scanned
        self.warning = warning

class DocumentExtractionResult:
    def __init__(self):
        self.pages: List[PageExtractionResult] = []
        self.total_pages: int = 0
        self.warnings: List[str] = []
        self.detected_tables: int = 0
        self.scanned_page_count: int = 0

class DocumentExtractor:

    @staticmethod
    def extract_pdf(file_bytes: bytes) -> DocumentExtractionResult:
        result = DocumentExtractionResult()
        
        # Check password protection or corruption
        try:
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            if doc.is_encrypted:
                raise ValueError("Password-protected PDF files cannot be processed. Please provide an unencrypted document.")
            result.total_pages = len(doc)
            if result.total_pages == 0:
                raise ValueError("The provided PDF contains no readable pages.")
        except Exception as e:
            if "encrypted" in str(e).lower() or "password" in str(e).lower():
                raise ValueError("Password-protected PDF files cannot be processed. Please provide an unencrypted document.")
            raise ValueError(f"Corrupted or invalid PDF document: {str(e)}")

        # Optional pdfplumber instance for tabular data
        plumber_doc = None
        try:
            plumber_doc = pdfplumber.open(io.BytesIO(file_bytes))
        except Exception:
            pass

        for page_idx in range(result.total_pages):
            page_num = page_idx + 1
            try:
                page = doc.load_page(page_idx)
                text = page.get_text("text").strip()
                is_scanned = False
                page_warning = None

                # Table extraction enrichment via pdfplumber
                if plumber_doc and page_idx < len(plumber_doc.pages):
                    try:
                        plumber_page = plumber_doc.pages[page_idx]
                        tables = plumber_page.extract_tables()
                        if tables:
                            table_texts = []
                            for t in tables:
                                if t and len(t) > 0:
                                    result.detected_tables += 1
                                    rows = [" | ".join([cell if cell else "" for cell in row]) for row in t]
                                    table_texts.append("\n[Table Data]:\n" + "\n".join(rows))
                            if table_texts:
                                text += "\n" + "\n".join(table_texts)
                    except Exception as te:
                        logger.debug(f"Plumber table extraction notice for page {page_num}: {te}")

                # Check if page is scanned (less than 40 chars of text but has images)
                if len(text) < 40:
                    image_list = page.get_images(full=True)
                    if image_list:
                        is_scanned = True
                        result.scanned_page_count += 1
                        ocr_text = DocumentExtractor._run_ocr_on_page(page)
                        if ocr_text:
                            text = (text + "\n" + ocr_text).strip()
                        else:
                            page_warning = f"Page {page_num} appears to be a scanned image, but OCR produced minimal text."

                result.pages.append(PageExtractionResult(
                    page_number=page_num,
                    text=text,
                    is_scanned=is_scanned,
                    warning=page_warning
                ))
                if page_warning:
                    result.warnings.append(page_warning)

            except Exception as pe:
                warning_msg = f"Page {page_num} could not be read. Processing continued with the remaining pages."
                logger.warning(warning_msg)
                result.warnings.append(warning_msg)
                result.pages.append(PageExtractionResult(
                    page_number=page_num,
                    text="",
                    warning=warning_msg
                ))

        if plumber_doc:
            try:
                plumber_doc.close()
            except Exception:
                pass
        doc.close()

        return result

    @staticmethod
    def _run_ocr_on_page(page: pymupdf.Page) -> str:
        """Renders page to pixmap and executes Tesseract OCR with fallback."""
        try:
            pix = page.get_pixmap(dpi=150)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img)
            return text.strip()
        except Exception as e:
            logger.debug(f"Tesseract OCR unavailable or skipped: {e}")
            return ""

    @staticmethod
    def extract_docx(file_bytes: bytes) -> DocumentExtractionResult:
        result = DocumentExtractionResult()
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_lines.append(p.text.strip())
            
            for t in doc.tables:
                result.detected_tables += 1
                rows = []
                for row in t.rows:
                    rows.append(" | ".join([cell.text.strip() for cell in row.cells]))
                text_lines.append("\n[Table Data]:\n" + "\n".join(rows))

            full_text = "\n\n".join(text_lines)
            # Partition into simulated page boundaries (approx 2500 chars per page)
            pages = DocumentExtractor._split_text_into_pages(full_text)
            result.total_pages = len(pages)
            for idx, p_text in enumerate(pages):
                result.pages.append(PageExtractionResult(page_number=idx + 1, text=p_text))
            return result
        except Exception as e:
            raise ValueError(f"Failed to read DOCX document: {str(e)}")

    @staticmethod
    def extract_legacy_doc(file_bytes: bytes) -> DocumentExtractionResult:
        """Safe fallback reader for binary legacy .doc files without crashing."""
        result = DocumentExtractionResult()
        try:
            # Decode printable strings from binary stream safely
            text = file_bytes.decode("latin-1", errors="ignore")
            # Extract printable sequences of 4+ characters
            printable = re.findall(r'[A-Za-z0-9\s.,;:\'\"?!(){}\[\]/\\+\-=]{4,}', text)
            filtered = "\n".join([line.strip() for line in printable if len(line.strip()) > 10])
            pages = DocumentExtractor._split_text_into_pages(filtered)
            result.total_pages = max(1, len(pages))
            result.warnings.append("Document was in legacy DOC format. Text extracted via compatibility stream parser.")
            for idx, p_text in enumerate(pages):
                result.pages.append(PageExtractionResult(page_number=idx + 1, text=p_text))
            return result
        except Exception as e:
            raise ValueError(f"Legacy DOC file could not be parsed: {str(e)}")

    @staticmethod
    def extract_pptx(file_bytes: bytes) -> DocumentExtractionResult:
        result = DocumentExtractionResult()
        try:
            prs = Presentation(io.BytesIO(file_bytes))
            result.total_pages = len(prs.slides)
            for idx, slide in enumerate(prs.slides):
                slide_num = idx + 1
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_texts.append(paragraph.text.strip())
                    elif shape.has_table:
                        result.detected_tables += 1
                        t_rows = []
                        for row in shape.table.rows:
                            t_rows.append(" | ".join([c.text.strip() for c in row.cells]))
                        slide_texts.append("\n[Slide Table]:\n" + "\n".join(t_rows))
                result.pages.append(PageExtractionResult(
                    page_number=slide_num,
                    text="\n".join(slide_texts)
                ))
            return result
        except Exception as e:
            raise ValueError(f"Failed to extract PPTX presentation: {str(e)}")

    @staticmethod
    def extract_txt(file_bytes: bytes) -> DocumentExtractionResult:
        result = DocumentExtractionResult()
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")
            result.warnings.append("File decoded using Latin-1 fallback encoding.")

        pages = DocumentExtractor._split_text_into_pages(text)
        result.total_pages = len(pages)
        for idx, p_text in enumerate(pages):
            result.pages.append(PageExtractionResult(page_number=idx + 1, text=p_text))
        return result

    @staticmethod
    def extract_image(file_bytes: bytes, file_ext: str) -> DocumentExtractionResult:
        result = DocumentExtractionResult()
        result.total_pages = 1
        result.scanned_page_count = 1
        try:
            img = Image.open(io.BytesIO(file_bytes))
            text = pytesseract.image_to_string(img).strip()
            if not text:
                result.warnings.append("OCR detected image but no text could be extracted.")
            result.pages.append(PageExtractionResult(
                page_number=1,
                text=text,
                is_scanned=True
            ))
            return result
        except Exception as e:
            logger.info(f"Image OCR execution notice: {e}")
            result.pages.append(PageExtractionResult(
                page_number=1,
                text="",
                warning="OCR is not configured in this environment or image text was unreadable."
            ))
            result.warnings.append("Image uploaded. OCR processing skipped.")
            return result

    @staticmethod
    def _split_text_into_pages(text: str, chars_per_page: int = 2500) -> List[str]:
        if not text:
            return [""]
        paragraphs = text.split("\n\n")
        pages = []
        current_page = []
        current_len = 0
        for p in paragraphs:
            current_page.append(p)
            current_len += len(p)
            if current_len >= chars_per_page:
                pages.append("\n\n".join(current_page))
                current_page = []
                current_len = 0
        if current_page:
            pages.append("\n\n".join(current_page))
        return pages if pages else [text]
